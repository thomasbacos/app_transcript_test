"""
Reference documents ("grounding") for a transcription: read PDF / Word / PowerPoint / Excel / text
files, extract the vocabulary that speech-to-text tends to mishear (names, acronyms, product and
technical terms, figures), and keep a text excerpt for the AI correction pass.

    ctx = grounding.load(["deck.pptx", "notes.pdf"])
    ctx["terms"]    -> ["Acme", "Kubernetes", "ISO 20022", ...]   (sent to the transcription model)
    ctx["excerpt"]  -> text given to the correction model
"""
import os, re, collections

AUDIO_EXT = {".m4a", ".mp3", ".mp4", ".wav", ".mov", ".webm", ".mpeg", ".mpga", ".aac", ".ogg", ".flac",
             ".mkv", ".wma", ".m4v", ".avi"}
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".eml"}
MAX_EXCERPT = 40_000          # characters of document text handed to the correction model


def kind(path):
    ext = os.path.splitext(path)[1].lower()
    return "audio" if ext in AUDIO_EXT else ("doc" if ext in DOC_EXT else None)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages[:300])
        if ext == ".docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if ext == ".pptx":
            from pptx import Presentation
            parts = []
            for i, slide in enumerate(Presentation(path).slides, 1):
                parts.append("[slide %d]" % i)
                for shp in slide.shapes:
                    if shp.has_text_frame:
                        parts.append(shp.text_frame.text)
                    if getattr(shp, "has_table", False) and shp.has_table:
                        for row in shp.table.rows:
                            parts.append(" | ".join(c.text for c in row.cells))
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
            return "\n".join(parts)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:20]:
                parts.append("[sheet %s]" % ws.title)
                for n, row in enumerate(ws.iter_rows(values_only=True)):
                    if n > 400:
                        break
                    cells = [str(v) for v in row if v not in (None, "")]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        if ext in (".html", ".htm"):
            raw = open(path, encoding="utf-8", errors="ignore").read()
            return re.sub(r"<[^>]+>", " ", raw)
        return open(path, encoding="utf-8", errors="ignore").read()
    except Exception as e:
        return "[could not read %s: %s]" % (os.path.basename(path), e)


STOP = set("""a à au aux avec ce ces cette dans de des du elle en et il ils je la le les leur lui mais me même
mes moi mon ne nos notre nous on ou où par pas pour qu que qui sa se ses son sur ta te tes toi ton tu un une
vos votre vous the a an and are as at be by for from has have in is it its of on or that the to was were will
with this these those our we you your their they he she not but if then than so such can may also all any
le la les slide sheet page total yes no oui non ok""".split())


def terms_from(text, limit=90):
    """Vocabulary likely to be misheard: acronyms, CamelCase / mixed tokens, words with digits, and
    capitalised names or multi-word proper nouns; ranked by frequency."""
    cnt = collections.Counter()
    for m in re.finditer(r"\b[A-Z][A-Z0-9&\-]{1,9}s?\b", text):                    # acronyms: ISO, NASA, KPI
        cnt[m.group()] += 2
    for m in re.finditer(r"\b[A-Za-z]+(?:\s?\d[\d\.]*)+\b|\b\d+[A-Za-z]+\b", text):   # ISO 20022, T24, 5G
        if any(c.isalpha() for c in m.group()) and len(m.group()) <= 20:
            cnt[m.group().strip()] += 2
    for m in re.finditer(r"\b[a-z]+[A-Z][A-Za-z]+\b|\b[A-Z][a-z]+[A-Z][A-Za-z]*\b", text):  # CamelCase
        cnt[m.group()] += 2
    for m in re.finditer(r"(?<![\.\!\?]\s)(?<!^)\b([A-ZÀ-Ý][a-zà-ÿ]{2,}(?:[ \-][A-ZÀ-Ý][a-zà-ÿ]{2,}){0,3})\b",
                         text, re.M):                                              # Proper Names
        w = m.group(1)
        if w.lower() not in STOP:
            cnt[w] += 1
    out = []
    for term, n in cnt.most_common():
        if n < 2 and len(out) > 30:
            break
        if term.lower() in STOP or len(term) < 2 or any(term.lower() == o.lower() for o in out):
            continue
        out.append(term)
        if len(out) >= limit:
            break
    return out


def load(paths):
    docs, texts = [], []
    for p in paths:
        t = extract_text(p)
        t = re.sub(r"[ \t]+", " ", t)
        t = re.sub(r"\n{3,}", "\n\n", t).strip()
        docs.append({"path": p, "name": os.path.basename(p), "chars": len(t)})
        texts.append("### %s\n%s" % (os.path.basename(p), t))
    full = "\n\n".join(texts)
    excerpt = full if len(full) <= MAX_EXCERPT else full[:MAX_EXCERPT] + "\n[...truncated]"
    return {"docs": docs, "terms": terms_from(full), "excerpt": excerpt}
